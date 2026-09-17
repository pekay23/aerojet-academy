"""Extract text from answer key documents only (fast version)."""
from __future__ import annotations

from pathlib import Path
import re

AEROJET_ROOT = Path(
    r"C:\Users\Pekay\OneDrive - Ghana Communication Technology University\AerojetAviation"
)
OUT_DIR = Path(__file__).parent / "doc-text"
MODULES = ["M1", "M2", "M3", "M4", "M5", "M6", "M7", "M11A", "M13", "M17"]

ANSWER_KEYWORDS = re.compile(
    r"answer|answers|answer key|with answers|highlighted|updated|correct",
    re.IGNORECASE,
)

module_folders = {
    "M1": "Module 1 - Mathematics",
    "M2": "Module 2 - Physics",
    "M3": "Module 3 - Electrical Fundamentals",
    "M4": "Module 4 - Electronic Fundamentals",
    "M5": "Module 5 - Digital Techniques (Electronic Instrument Systems)",
    "M6": "Module 6 - Materials and Hardware",
    "M7": "Module 7 - Maintenance Practices",
    "M11A": "Module 11A - Turbine Aeroplane Aerodynamics",
    "M13": "Module 13 - Aircraft Aerodynamics, Structures & Systems",
    "M17": "Module 17 - Propellers",
}


def read_docx(path: Path) -> str:
    from docx import Document
    try:
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    lines.append(f"[TABLE] {row_text}")
        return "\n".join(lines)
    except Exception as e:
        return f"[DOCX ERROR: {e}]"


def read_pdf(path: Path) -> str:
    import pdfplumber
    pages = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            for i, page in enumerate(pdf.pages):
                try:
                    text = page.extract_text()
                    if text:
                        pages.append(f"--- Page {i+1} ---\n{text}")
                except Exception:
                    pass
    except Exception as e:
        return f"[PDF ERROR: {e}]"
    return "\n\n".join(pages)


def extract_answer_keys():
    total = 0
    for module, folder_name in module_folders.items():
        folder = AEROJET_ROOT / folder_name
        if not folder.exists():
            print(f"  SKIP {module}: folder not found")
            continue

        out_module = OUT_DIR / module
        out_module.mkdir(parents=True, exist_ok=True)

        for path in folder.rglob("*"):
            if not path.is_file():
                continue
            if path.suffix.lower() not in {".pdf", ".docx"}:
                continue
            if not ANSWER_KEYWORDS.search(path.name):
                continue

            out_path = out_module / f"{path.stem}.txt"
            if out_path.exists():
                continue

            print(f"  Extracting {module}: {path.name}")
            if path.suffix.lower() == ".pdf":
                text = read_pdf(path)
            else:
                text = read_docx(path)
            out_path.write_text(text, encoding="utf-8")
            total += 1

    print(f"\nDone. Extracted {total} answer key documents.")


if __name__ == "__main__":
    extract_answer_keys()
