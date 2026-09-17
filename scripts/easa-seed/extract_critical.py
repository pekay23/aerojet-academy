"""Extract ONLY critical answer key documents for each module."""
from __future__ import annotations

from pathlib import Path
import pdfplumber
from docx import Document

AEROJET_ROOT = Path(
    r"C:\Users\Pekay\OneDrive - Ghana Communication Technology University\AerojetAviation"
)
OUT_DIR = Path(__file__).parent / "doc-text"
MODULES = ["M1", "M2", "M3", "M4", "M5", "M6", "M7", "M11A", "M13", "M17"]

CRITICAL_DOCS = {
    "M1": [],  # Already extracted
    "M2": [
        "Module 2 - Physics/Questions/ANSWERS MODULE 02 QUESTIONS BANK BySFoS 2.pdf",
        "Module 2 - Physics/Questions/AEROJET M2 SAMPLE SET 10 WITH ANSWERS R.pdf",
        "Module 2 - Physics/Questions/THRUST2 (ACIDS AND ALKALIS) with answers.docx",
    ],
    "M3": [
        "Module 3 - Electrical Fundamentals/Aviators Question1.docx",
    ],
    "M4": [
        "Module 4 - Electronic Fundamentals/M4 Standard EASA Exam.docx",
    ],
    "M5": [
        "Module 5 - Digital Techniques (Electronic Instrument Systems)/Questions/M5 Standard EASA Exam.docx",
    ],
    "M6": [
        "Module 6 - Materials and Hardware/M6 Questions.docx",
    ],
    "M7": [
        "Module 7 - Maintenance Practices/M7 highlighted.docx",
        "Module 7 - Maintenance Practices/Updated.docx",
    ],
    "M11A": [],
    "M13": [
        "Module 13 - Aircraft Aerodynamics, Structures & Systems/exams q.pdf",
    ],
    "M17": [
        "Module 17 - Propellers/M17 - Propeller standard test.docx",
    ],
}


def read_docx(path: Path) -> str:
    try:
        doc = Document(str(path))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    lines.append(f"[TABLE] {row_text}")
        return "\n".join(lines)
    except Exception as e:
        return f"[DOCX ERROR: {e}]"


def read_pdf(path: Path) -> str:
    pages = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            for i, page in enumerate(pdf.pages):
                try:
                    text = page.extract_text()
                    if text:
                        pages.append(f"--- Page {i+1} ---\n{text}")
                except Exception:
                    pass
    except Exception as e:
        return f"[PDF ERROR: {e}]"
    return "\n\n".join(pages)


def extract_critical():
    for module, rel_paths in CRITICAL_DOCS.items():
        out_module = OUT_DIR / module
        out_module.mkdir(parents=True, exist_ok=True)
        for rel_path in rel_paths:
            path = AEROJET_ROOT / rel_path
            if not path.exists():
                print(f"  SKIP {module}: {path.name} not found")
                continue
            out_path = out_module / f"{path.stem}.txt"
            if out_path.exists():
                print(f"  SKIP {module}: {out_path.name} already exists")
                continue
            print(f"  Extracting {module}: {path.name}")
            if path.suffix.lower() == ".pdf":
                text = read_pdf(path)
            else:
                text = read_docx(path)
            out_path.write_text(text, encoding="utf-8")


if __name__ == "__main__":
    extract_critical()
    print("Done.")
