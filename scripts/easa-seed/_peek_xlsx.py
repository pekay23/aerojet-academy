"""Read xlsx sheet content."""
import sys
from openpyxl import load_workbook

path = sys.argv[1]
wb = load_workbook(path, data_only=True)
for sheet_name in wb.sheetnames:
    print(f"=== {sheet_name} ===")
    ws = wb[sheet_name]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if any(c is not None and str(c).strip() for c in row):
            print(" | ".join("" if c is None else str(c) for c in row))
        if i > 80:
            break
    print()
    if len(wb.sheetnames) == 1:
        break
